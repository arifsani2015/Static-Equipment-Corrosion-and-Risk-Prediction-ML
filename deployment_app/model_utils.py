"""
Core logic for the deployment app: loading checkpoints from Experiment 1
(corrosion-rate regression) and Experiment 2 (risk classification), model
selection, metrics, predictions, corrosion-rate reliability, SHAP and Excel
input/output.

The prediction logic intentionally mirrors the notebooks (Experiment 1/2/3)
so numbers produced here match the notebooks.

The feature schema is NEVER hard-coded: numeric/categorical columns are always
read from the fitted ColumnTransformer stored in (or rebuilt from) each
checkpoint -- see extract_schema(). Retrain with different features and the
app adapts automatically.
"""
import io
import json
import os
import pickle
import re
import tempfile
import urllib.parse
import urllib.request
import warnings

import numpy as np
import pandas as pd
import shap

RISK_ORDER = ["Low", "Medium", "Med-High", "High"]
RISK_COLORS = {"Low": "#4CAF50", "Medium": "#F1C40F", "Med-High": "#E67E22", "High": "#E74C3C"}
CONF_COLORS = {"High": "#4CAF50", "Medium": "#F1C40F", "Low": "#E74C3C"}

# Default display labels for common raw column names (only a helper -- any
# column not listed keeps its raw name until relabelled on the Settings page).
DEFAULT_DISPLAY = {
    "op_pressure": "Operating Pressure (Psig)", "op_temp": "Operating Temperature (F)",
    "inner_dia": "Inner Diameter (mm)", "nom_thk": "Nominal Thickness (mm)",
    "act_used_thk": "Actual Used Thickness (mm)", "min_req_thk": "Minimum Required Thickness (mm)",
    "n_insp": "Number of Inspections", "insp_age": "Inspection Age (Years)",
    "volume": "Volume (m3)", "allow_stress": "Allowable Stress (Psig)",
    "material": "Material", "equip_type": "Equipment Type", "insulation_exist": "Insulation",
    "phase": "Fluid Release Phase", "majority_phase": "Majority Fluid Phase",
    "selected_subsystem": "Subsystem", "service_code": "Service Code",
    "corrosion_agent": "Corrosion Agent", "cladding": "Cladding", "fluid": "Representative Fluid",
    "clscc": "ClSCC Susceptibility", "cr": "Corrosion Rate (mm/year)",
}

# Experiment 1 DISPLAY dict (copied from notebook cell A4), inverted
# (pretty label -> raw column). Only used for Experiment 1 checkpoints that do
# not store the preprocessor: their feature names are pretty labels.
P1_DISPLAY_INV = {
    "Operating Pressure (Psig)": "op_pressure", "Operating Temperature (F)": "op_temp",
    "Inner Diameter": "inner_dia", "Nominal Thickness": "nom_thk",
    "Material": "material", "Equipment Type": "equip_type", "Insulation": "insulation_exist",
    "Majority Fluid Phase": "majority_phase", "Subsystem": "selected_subsystem",
    "Service Code": "service_code", "Corrosion Agent": "corrosion_agent", "Cladding": "cladding",
}

# In Experiment 1 'phase' comes from Excel column "Selected Majority Fluid
# Phase"; in Experiment 2 'phase' comes from "Fluid Release Phase". Different
# quantities with the same raw name -> regression 'phase' is exposed to the
# user as 'majority_phase' and renamed back just before the preprocessor.
REG_COL_ALIAS_MODEL_TO_APP = {"phase": "majority_phase"}

try:
    from remote_checkpoints import REMOTE_CHECKPOINTS
except Exception:
    REMOTE_CHECKPOINTS = []

GITHUB_RELEASE_RE = re.compile(r"^https://github\.com/(?P<owner>[^/]+)/(?P<repo>[^/]+)/releases/download/"
                               r"(?P<tag>[^/]+)/(?P<name>[^/?#]+)")

CKPT_REG_RE = re.compile(r"checkpoint_reg_(?P<tag>.+)\.pkl$", re.I)
CKPT_P2_RE = re.compile(r"checkpoint_p2_(?P<source>measured|selected|estimated)_(?P<tag>.+)\.pkl$", re.I)

# App folder (where app.py lives), NOT the working directory.
APP_DIR = os.path.dirname(os.path.abspath(__file__))
# Extra folders to scan for checkpoints, taken from the CHECKPOINT_DIRS environment
# variable (separated by os.pathsep, e.g. "/data/models:/mnt/outputs"). Empty by default;
# useful when the .pkl files are mounted outside the app folder.
EXTRA_CHECKPOINT_DIRS = tuple(d.strip() for d in os.environ.get("CHECKPOINT_DIRS", "").split(os.pathsep)
                              if d.strip())
DEFAULT_SEARCH_DIRS = EXTRA_CHECKPOINT_DIRS + (APP_DIR,)

REG_METRIC_ORDER = ["R2", "RMSE", "MAE", "IA", "R2_CV", "RMSE_CV", "MAE_CV", "IA_CV"]
CLF_METRIC_ORDER = ["Accuracy", "F1_macro", "F1_weighted", "Kappa", "ROC_AUC",
                    "F1_macro_CV", "Accuracy_CV", "F1_weighted_CV", "Kappa_CV", "ROC_AUC_CV"]


# =====================================================================
# Display helpers
# =====================================================================
def resolve_display(col, overrides=None):
    """Display label: Settings-page override > default dictionary > raw name."""
    overrides = overrides or {}
    return overrides.get(col) or DEFAULT_DISPLAY.get(col, col)


def base_column(name):
    """Map a one-hot feature name ('Material = SS316L') back to 'Material'."""
    return name.split(" = ")[0]


# =====================================================================
# Schema & preprocessing
# =====================================================================
def extract_schema(res):
    """(numeric_cols, categorical_cols) read from the fitted ColumnTransformer
    ('pre') -- the single source of truth for the input schema."""
    pre = res.get("pre")
    if pre is None or not hasattr(pre, "transformers_"):
        raise ValueError("This checkpoint has no valid preprocessor ('pre').")
    num_cols, cat_cols = [], []
    for name, _t, cols in pre.transformers_:
        if name == "num":
            num_cols = list(cols)
        elif name == "cat":
            cat_cols = list(cols)
    if not num_cols and not cat_cols:
        raise ValueError("Could not read the column schema from the preprocessor.")
    alias = res.get("_col_alias_model_to_app") or {}
    if alias:
        num_cols = [alias.get(c, c) for c in num_cols]
        cat_cols = [alias.get(c, c) for c in cat_cols]
    return num_cols, cat_cols


def _to_model_frame(res, X):
    """App column names -> model column names, plus the same categorical
    normalisation as load_data() in the notebooks (astype(str).strip()), so
    Excel values such as 'CS ' or the number 1 still match learned categories.
    NaN stays NaN (filled by the model's imputer)."""
    alias = res.get("_col_alias_model_to_app") or {}
    X = X.rename(columns={v: k for k, v in alias.items()}).copy()
    for name, _t, cols in res["pre"].transformers_:
        for c in cols:
            if c not in X.columns:
                continue
            if name == "num":
                X[c] = pd.to_numeric(X[c], errors="coerce")
            elif name == "cat":
                col = X[c]
                mask = col.notna()
                X[c] = col.astype(object)
                X.loc[mask, c] = col[mask].astype(str).str.strip()
    return X


def _transform(res, X_app):
    """App columns -> feature matrix named/ordered like res['feat_names']."""
    X = _to_model_frame(res, X_app)
    Xt = res["pre"].transform(X)
    Xt = Xt.toarray() if hasattr(Xt, "toarray") else Xt
    gen = res.get("_pre_feat_names")
    if gen is not None:
        return pd.DataFrame(Xt, columns=gen).reindex(columns=res["feat_names"], fill_value=0.0)
    return pd.DataFrame(Xt, columns=res["feat_names"])


def rebuild_reg_preprocessor(res):
    """Experiment 1 checkpoints (run_regression, cell B3) did NOT store 'pre',
    only Xtr_t (training matrix AFTER preprocessing, pretty column names) and
    feat_names. Rebuild 'pre' EXACTLY from Xtr_t without retraining:
      - numeric: median of the Xtr_t column == original median (filling NaN
        with the median does not move the median),
      - categorical: categories read from one-hot names ('Material = CS');
        the mode of the decoded column == original mode (same argument).
    Verified: generated feature names must equal feat_names."""
    from sklearn.compose import ColumnTransformer
    from sklearn.impute import SimpleImputer
    from sklearn.pipeline import Pipeline
    from sklearn.preprocessing import OneHotEncoder

    feat_names = list(res["feat_names"])
    Xtr_t = res.get("Xtr_t")
    if Xtr_t is None:
        raise ValueError("This regression checkpoint stores neither 'pre' nor 'Xtr_t'.")
    Xtr_t = pd.DataFrame(np.asarray(Xtr_t), columns=feat_names)

    num_pretty = [n for n in feat_names if " = " not in n]
    cat_groups = {}
    for n in feat_names:
        if " = " in n:
            base, val = n.split(" = ", 1)
            cat_groups.setdefault(base, []).append((n, val))

    raw = pd.DataFrame(index=Xtr_t.index)
    num_raw, cat_raw, pretty_of = [], [], {}
    for n in num_pretty:
        r = P1_DISPLAY_INV.get(n, n)
        raw[r] = Xtr_t[n].astype(float).values
        num_raw.append(r); pretty_of[r] = n
    for base, items in cat_groups.items():
        r = P1_DISPLAY_INV.get(base, base)
        block = Xtr_t[[n for n, _ in items]].values
        vals = np.array([v for _, v in items], dtype=object)
        col = pd.Series(vals[block.argmax(axis=1)], index=Xtr_t.index, dtype=object)
        col[block.max(axis=1) <= 0] = np.nan
        raw[r] = col
        cat_raw.append(r); pretty_of[r] = base

    pre = ColumnTransformer([
        ("num", SimpleImputer(strategy="median"), num_raw),
        ("cat", Pipeline([("imp", SimpleImputer(strategy="most_frequent")),
                          ("oh", OneHotEncoder(handle_unknown="ignore"))]), cat_raw),
    ]).fit(raw)

    gen = [pretty_of[c] for c in num_raw]
    oh = pre.named_transformers_["cat"].named_steps["oh"]
    for c, cats in zip(cat_raw, oh.categories_):
        gen += [f"{pretty_of[c]} = {v}" for v in cats]
    if set(gen) != set(feat_names) or len(gen) != len(feat_names):
        raise ValueError("Rebuilt regression preprocessor does not match the checkpoint feat_names.")
    return pre, gen


def prepare_reg_store(store):
    """Run once per loaded regression checkpoint (idempotent): rebuild 'pre'
    if missing; alias 'phase' -> 'majority_phase' if the stored 'pre' uses it."""
    for _key, res in store.items():
        if res.get("_prepared"):
            continue
        if res.get("pre") is None:
            pre, gen = rebuild_reg_preprocessor(res)
            res["pre"], res["_pre_feat_names"], res["_pre_reconstructed"] = pre, gen, True
        else:
            cols = [c for _n, _t, cc in res["pre"].transformers_ for c in list(cc)]
            alias = {k: v for k, v in REG_COL_ALIAS_MODEL_TO_APP.items() if k in cols}
            if alias:
                res["_col_alias_model_to_app"] = alias
        res["_prepared"] = True
    return store


def get_training_defaults(res):
    """{app_column: value the model's imputer would use} -- training median
    (numeric) and mode (categorical). Used to pre-fill the manual forms."""
    alias = res.get("_col_alias_model_to_app") or {}
    out = {}
    try:
        for name, trans, cols in res["pre"].transformers_:
            imp = trans if name == "num" else getattr(trans, "named_steps", {}).get("imp")
            if imp is None or not hasattr(imp, "statistics_"):
                continue
            for c, v in zip(cols, imp.statistics_):
                out[alias.get(c, c)] = float(v) if name == "num" else str(v)
    except Exception:
        pass
    return out


def get_categories_from_checkpoint(res, cat_cols):
    """Valid categories per categorical column, read from the fitted OneHotEncoder."""
    try:
        _n, all_cat = extract_schema(res)
        oh = res["pre"].named_transformers_["cat"].named_steps["oh"]
        full = {col: [str(v) for v in cats] for col, cats in zip(all_cat, oh.categories_)}
        return {c: full.get(c, []) for c in cat_cols}
    except Exception:
        return {c: [] for c in cat_cols}


def get_effective_categories(res, cat_cols, category_overrides=None):
    """Auto-detected categories + extra categories added on the Settings page."""
    auto = get_categories_from_checkpoint(res, cat_cols)
    overrides = category_overrides or {}
    merged = {}
    for c in cat_cols:
        combined = list(auto.get(c, []))
        for v in overrides.get(c, []):
            if v not in combined:
                combined.append(v)
        merged[c] = combined
    return merged


def get_numeric_ranges(res, num_cols):
    """(min, max) of each numeric feature as seen during training. Xtr_t
    columns are pretty labels, so values are read by POSITION (numeric
    features always come first, in num_cols order)."""
    Xtr_t = res.get("Xtr_t")
    if Xtr_t is None:
        return {}
    arr = pd.DataFrame(np.asarray(Xtr_t))
    all_num, _ = extract_schema(res)
    ranges = {}
    for i, c in enumerate(all_num):
        if c in num_cols and i < arr.shape[1]:
            col = pd.to_numeric(arr.iloc[:, i], errors="coerce")
            ranges[c] = (float(col.min()), float(col.max()))
    return ranges


# =====================================================================
# Loading & discovery
# =====================================================================
def _classify(obj):
    if isinstance(obj, dict) and "table" in obj:
        return "classification", obj
    if isinstance(obj, dict) and obj and all(isinstance(v, dict) and "table" in v for v in obj.values()):
        return "regression", obj
    raise ValueError("Unrecognised checkpoint format -- expected checkpoint_reg_*.pkl "
                     "(Experiment 1) or checkpoint_p2_*.pkl (Experiment 2).")


def load_checkpoint(source):
    """Load ONE checkpoint and detect its type:
      - Experiment 1 (regression): nested dict {scenario: {...}}
      - Experiment 2 (classification): flat dict containing 'table'.
    source: bytes, or a binary file object (preferred for big files: it is
    unpickled straight from the stream, without a second full copy in RAM)."""
    obj = pickle.loads(source) if isinstance(source, (bytes, bytearray)) else pickle.load(source)
    return _classify(obj)


_SKLEARN_MISMATCH_RE = re.compile(r"from version (\d+(?:\.\d+)+\w*) when using version (\d+(?:\.\d+)+\w*)")


def load_checkpoint_verbose(source):
    """load_checkpoint() + library-version warnings raised while unpickling
    (scikit-learn InconsistentVersionWarning, XGBoost serialisation notes...).
    Returns (kind, obj, warnings)."""
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        kind, obj = load_checkpoint(source)
    msgs, groups = [], {}
    for w in caught:
        text = str(w.message)
        m = _SKLEARN_MISMATCH_RE.search(text)
        if m:  # scikit-learn warns once per estimator object: merge them into ONE line
            groups[m.groups()] = groups.get(m.groups(), 0) + 1
            continue
        if ("version" in text.lower() or "InconsistentVersion" in type(w.message).__name__
                or "serialized" in text.lower()):
            short = text.split("\n")[0][:300]
            if short not in msgs:
                msgs.append(short)
    msgs = [f"scikit-learn: trained with {a}, running {b}."
            for (a, b), n in groups.items()] + msgs
    if kind == "regression":
        prepare_reg_store(obj)
    return kind, obj, msgs


def sklearn_version_in_pickle(source, max_ops=400_000):
    """scikit-learn version a checkpoint was TRAINED with, read from the file
    itself. pickletools only parses the opcodes -- nothing is executed and
    scikit-learn does not even have to be importable -- and stops at the first
    estimator (its state ends with '_sklearn_version'), so it is fast even for
    files of several hundred MB. source: a path, or a binary file object.
    Returns e.g. '1.5.2', or None when it cannot be found (other formats)."""
    import pickletools
    opened = None
    try:
        if isinstance(source, (str, os.PathLike)):
            opened = open(source, "rb")
            stream = opened
        else:
            stream = source
            if hasattr(stream, "seek"):
                stream.seek(0)
        after_key = False
        for n, (_op, arg, _pos) in enumerate(pickletools.genops(stream)):
            if n > max_ops:
                break
            if isinstance(arg, str):
                if after_key:
                    return arg if re.fullmatch(r"\d+(\.\d+)+\w*", arg) else None
                after_key = (arg == "_sklearn_version")
        return None
    except Exception:
        return None
    finally:
        if opened is not None:
            opened.close()
        elif hasattr(source, "seek"):
            try:
                source.seek(0)
            except Exception:
                pass


def record_trained_version(st, kind, source):
    """Remember, per checkpoint kind, the scikit-learn version it was trained
    with (shown on the Settings page). Never raises."""
    try:
        v = sklearn_version_in_pickle(source) if source is not None else None
    except Exception:
        v = None
    st.session_state.setdefault("_trained_sklearn", {})[kind] = v


def find_file(base_dir, must_contain=(), must_end=(), must_not=()):
    """Same logic as find_file() in Experiment 3: first file under base_dir
    matching the name pattern, or None. Never raises."""
    if not base_dir or not os.path.isdir(base_dir):
        return None
    for root, _, files in os.walk(base_dir, followlinks=True):
        for f in sorted(files):
            low = f.lower()
            if (all(s in low for s in must_contain) and low.endswith(must_end)
                    and not any(s in low for s in must_not)):
                return os.path.join(root, f)
    return None


def discover_checkpoint_candidates(base_dirs=DEFAULT_SEARCH_DIRS):
    """Scan for ALL Experiment 1 & 2 checkpoints under base_dirs (any depth).
    Only file NAMES are read; identical files in two places count once."""
    found, seen = [], set()
    for base in base_dirs:
        if not base or not os.path.isdir(base):
            continue
        for root, _dirs, files in os.walk(base, followlinks=True):
            for f in sorted(files):
                m_reg, m_p2 = CKPT_REG_RE.search(f), CKPT_P2_RE.search(f)
                if not (m_reg or m_p2):
                    continue
                path = os.path.join(root, f)
                real = os.path.realpath(path)
                try:
                    size, mtime = os.path.getsize(real), os.path.getmtime(real)
                except OSError:
                    continue
                ident = (f.lower(), size)
                if real in seen or ident in seen:
                    continue
                seen.update({real, ident})
                if m_p2:
                    found.append(dict(kind="classification", path=path, name=f, tag=m_p2.group("tag"),
                                      source=m_p2.group("source").lower(), size=size, mtime=mtime))
                else:
                    found.append(dict(kind="regression", path=path, name=f, tag=m_reg.group("tag"),
                                      source=None, size=size, mtime=mtime))
    return found


def candidate_label(c):
    extra = f" | CR {c['source']}" if c.get("source") else ""
    return f"{c['name']}  (TAG {c['tag']}{extra}, {c['size'] / 1e6:.1f} MB) -- {os.path.dirname(c['path'])}"


def choose_default_pair(cands):
    """Default pair: same TAG for regression & classification, classification
    CR type 'selected', otherwise the newest files."""
    regs = sorted([c for c in cands if c["kind"] == "regression"], key=lambda c: -c["mtime"])
    clfs = sorted([c for c in cands if c["kind"] == "classification"],
                  key=lambda c: (c["source"] != "selected", -c["mtime"]))
    for r in regs:
        for k in clfs:
            if k["tag"] == r["tag"]:
                return r, k
    return (regs[0] if regs else None), (clfs[0] if clfs else None)


def _load_path(path, _mtime, _size):
    with open(path, "rb") as f:
        return load_checkpoint_verbose(f)


try:  # one shared copy per file for ALL sessions (Streamlit Cloud has ~2.7 GB RAM per app)
    import streamlit as _st
    _load_path_cached = _st.cache_resource(show_spinner=False, max_entries=6)(_load_path)
except Exception:  # pragma: no cover - streamlit always installed with this app
    _load_path_cached = _load_path


def load_candidate(cand):
    """Load a checkpoint found on disk. Cached by (path, modified time, size):
    every visitor of a hosted app shares the same object instead of loading
    its own copy, and a replaced file is reloaded automatically."""
    path = cand["path"]
    return _load_path_cached(path, os.path.getmtime(path), os.path.getsize(path))


def reg_scenario_cr_type(key):
    """'Selected_log1p' -> 'selected'; 'Measured_raw' -> 'measured'."""
    return str(key).split("_")[0].lower()


def clf_source_from(res, filename=None):
    """CR type used to train the classifier (from file name or res['scenario'])."""
    if filename:
        m = CKPT_P2_RE.search(os.path.basename(str(filename)))
        if m:
            return m.group("source").lower()
    sc = str((res or {}).get("scenario", "")).lower()
    for t in ("selected", "measured", "estimated"):
        if t in sc:
            return t
    return None


# =====================================================================
# Models, metrics & quality
# =====================================================================
def model_names(res):
    """Fitted models, ordered as in the checkpoint's results table (best first)."""
    fitted = list(res.get("fitted", {}).keys())
    tbl = res.get("table")
    order = list(tbl["Model"]) if tbl is not None and "Model" in tbl.columns else []
    return [m for m in order if m in fitted] + [m for m in fitted if m not in order]


def resolve_model(res, model_name=None):
    return model_name if model_name in res.get("fitted", {}) else res["best_name"]


def metrics_row(res, model_name=None):
    """Test + CV metrics of one model. Regression tables already contain CV
    columns; classification CV metrics live in table_default_cv /
    table_optuna_cv and are merged according to overall_source."""
    name = resolve_model(res, model_name)
    tbl = res.get("table")
    if tbl is None or "Model" not in tbl.columns:
        return None
    row = tbl[tbl["Model"] == name]
    if not len(row):
        return None
    row = row.iloc[0].to_dict()
    cv_tbl = res.get("table_optuna_cv") if res.get("overall_source") == "Optuna" else res.get("table_default_cv")
    if cv_tbl is not None and len(cv_tbl) and "Model" in cv_tbl.columns:
        cv_row = cv_tbl[cv_tbl["Model"] == name]
        if len(cv_row):
            for k, v in cv_row.iloc[0].to_dict().items():
                row.setdefault(k, v)
    return row


def get_model_metrics_row(res, model_name=None):  # backwards-compatible name
    return metrics_row(res, model_name)


def train_score(res, model_name, kind):
    """Training-set score used for the overfitting check (R2 or F1_macro)."""
    name = resolve_model(res, model_name)
    if kind == "regression":
        v = (res.get("train_r2") or {}).get(name)
        return float(v) if v is not None else None
    tt = res.get("table_train")
    if tt is not None and "Model" in tt.columns and "F1_macro" in tt.columns:
        r = tt[tt["Model"] == name]
        if len(r):
            return float(r["F1_macro"].iloc[0])
    return None


def quality_label(row, kind, train=None):
    """Rule-of-thumb quality rating shown on the Settings page.
    Regression: R2_CV (or R2) >= 0.75 Good, >= 0.50 Fair, else Weak.
    Classification: F1_macro_CV (or F1_macro) >= 0.75 Good, >= 0.60 Fair, else Weak.
    Adds an overfitting note when train score exceeds test score by > 0.20."""
    if not row:
        return "Unknown", ""
    if kind == "regression":
        key = "R2_CV" if pd.notna(row.get("R2_CV", np.nan)) else "R2"
        test_key, good, fair = "R2", 0.75, 0.50
    else:
        key = "F1_macro_CV" if pd.notna(row.get("F1_macro_CV", np.nan)) else "F1_macro"
        test_key, good, fair = "F1_macro", 0.75, 0.60
    v = row.get(key)
    if v is None or pd.isna(v):
        return "Unknown", ""
    label = "Good" if v >= good else ("Fair" if v >= fair else "Weak")
    note = f"based on {key} = {v:.3f}"
    t = row.get(test_key)
    if train is not None and t is not None and pd.notna(t) and train - t > 0.20:
        note += f"; possible overfitting (train {train:.2f} vs test {t:.2f})"
    return label, note


def model_metrics_table(res, kind):
    """All models of one checkpoint/scenario with metrics + quality label."""
    order = REG_METRIC_ORDER if kind == "regression" else CLF_METRIC_ORDER
    rows = []
    for m in model_names(res):
        r = metrics_row(res, m) or {"Model": m}
        q, note = quality_label(r, kind, train_score(res, m, kind))
        out = {"Model": m + (" ★ best" if m == res["best_name"] else "")}
        for k in order:
            if k in r and pd.notna(r[k]):
                out[k] = float(r[k])
        out["Quality"] = q
        out["Notes"] = note
        rows.append(out)
    return pd.DataFrame(rows)


def scenario_summary(store):
    """Best model of every regression scenario (CR type x transform)."""
    rows = []
    for key, res in store.items():
        r = metrics_row(res) or {}
        q, _ = quality_label(r, "regression", train_score(res, None, "regression"))
        rows.append({"Scenario": key, "CR type": reg_scenario_cr_type(key).capitalize(),
                     "Transform": "log1p" if res.get("use_log") else "raw",
                     "Best model": res["best_name"],
                     **{k: float(r[k]) for k in ("R2", "RMSE", "R2_CV", "RMSE_CV") if k in r and pd.notna(r[k])},
                     "Quality": q})
    return pd.DataFrame(rows)


def pick_default_scenario(store, prefer_cr_type=None):
    """Default scenario = lowest RMSE_CV (same criterion as C1 in Experiment
    1), preferring scenarios whose CR type matches the classifier."""
    keys = list(store.keys())
    if not keys:
        return None
    pool = [k for k in keys if prefer_cr_type and reg_scenario_cr_type(k) == prefer_cr_type] or keys

    def score(k):
        v = (metrics_row(store[k]) or {}).get("RMSE_CV")
        return float(v) if v is not None and pd.notna(v) else float("inf")
    return min(pool, key=score)


# =====================================================================
# Session state
# =====================================================================
def init_session_state(st):
    defaults = {
        "reg_store": None, "reg_scenario": None, "reg_model": None, "reg_path": None, "reg_label": None,
        "clf_res": None, "clf_model": None, "clf_path": None, "clf_label": None, "clf_source": None,
        "display_overrides": {},   # {raw column: display label}
        "header_overrides": {},    # {raw column: Excel header used in template/batch}
        "default_overrides": {},   # {raw column: default value in forms / for hidden columns}
        "category_overrides": {},  # {raw column: [extra categories]}
        "excluded_cols": set(),    # columns hidden from forms and templates
        "_auto_discover_tried": False, "_auto_discover_msgs": [],
        "_ckpt_candidates": [], "_load_warnings": [], "_trained_sklearn": {},
        "hist_cr": [], "hist_risk": [], "hist_combo": [],
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v


def set_reg_scenario(st, key):
    """Switch regression scenario; model resets to that scenario's best."""
    st.session_state["reg_scenario"] = key
    st.session_state["reg_model"] = st.session_state["reg_store"][key]["best_name"]


def activate_reg(st, store, path=None, label=None, prefer_cr_type=None, version_source=None):
    record_trained_version(st, "regression", version_source if version_source is not None else path)
    st.session_state["reg_store"] = store
    st.session_state["reg_path"] = path
    st.session_state["reg_label"] = label
    ctype = prefer_cr_type or st.session_state.get("clf_source")
    set_reg_scenario(st, pick_default_scenario(store, ctype))


def activate_clf(st, res, path=None, label=None, source=None, version_source=None):
    record_trained_version(st, "classification", version_source if version_source is not None else path)
    st.session_state["clf_res"] = res
    st.session_state["clf_model"] = res["best_name"]
    st.session_state["clf_path"] = path
    st.session_state["clf_label"] = label
    st.session_state["clf_source"] = source or clf_source_from(res, label)


def active_reg(st):
    """(res, model_name) of the active regression model, or (None, None)."""
    store = st.session_state.get("reg_store")
    if not store:
        return None, None
    res = store[st.session_state["reg_scenario"]]
    return res, resolve_model(res, st.session_state.get("reg_model"))


def active_clf(st):
    res = st.session_state.get("clf_res")
    if res is None:
        return None, None
    return res, resolve_model(res, st.session_state.get("clf_model"))


# =====================================================================
# Remote checkpoints (GitHub Release) -- for .pkl files too large for git
# =====================================================================
class _DropAuthOnRedirect(urllib.request.HTTPRedirectHandler):
    """GitHub redirects release downloads to a storage host with its own
    signed URL; forwarding the GitHub token there makes the download fail,
    so the Authorization header is dropped when the host changes."""

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        new = super().redirect_request(req, fp, code, msg, headers, newurl)
        if new is not None and urllib.parse.urlparse(newurl).netloc != urllib.parse.urlparse(req.full_url).netloc:
            new.remove_header("Authorization")
        return new


_OPENER = urllib.request.build_opener(_DropAuthOnRedirect)
GITHUB_API = "https://api.github.com"
_UA = {"User-Agent": "corrosion-risk-streamlit-app"}


def _github_token():
    """Token for PRIVATE repositories: Streamlit Cloud 'Secrets'
    (GITHUB_TOKEN = "...") or the GITHUB_TOKEN environment variable.
    Never stored in the code or the repository."""
    tok = os.environ.get("GITHUB_TOKEN")
    if tok:
        return tok.strip()
    try:
        import streamlit as st
        tok = st.secrets.get("GITHUB_TOKEN")
        return str(tok).strip() if tok else None
    except Exception:
        return None


def _open_remote(url, token=None, use_api=None):
    """Open a download stream. GitHub Release links of PRIVATE repositories
    are only downloadable through the GitHub API with a token, so when a
    token is available (or use_api=True) the asset is resolved via the API;
    otherwise the public link is used directly."""
    m = GITHUB_RELEASE_RE.match(url)
    if m and (token or use_api):
        auth = {"Authorization": f"Bearer {token}"} if token else {}
        api = (f"{GITHUB_API}/repos/{m['owner']}/{m['repo']}/releases/tags/"
               f"{urllib.parse.quote(urllib.parse.unquote(m['tag']))}")
        with _OPENER.open(urllib.request.Request(api, headers={**_UA, **auth,
                          "Accept": "application/vnd.github+json"}), timeout=30) as r:
            release = json.load(r)
        name = urllib.parse.unquote(m["name"])
        asset = next((a for a in release.get("assets", []) if a.get("name") == name), None)
        if asset is None:
            raise FileNotFoundError(f"release '{m['tag']}' has no file named '{name}'")
        return _OPENER.open(urllib.request.Request(asset["url"], headers={**_UA, **auth,
                            "Accept": "application/octet-stream"}), timeout=60)
    return _OPENER.open(urllib.request.Request(url, headers=_UA), timeout=60)


def _download_dir():
    """checkpoints/ next to app.py; falls back to a temp folder if that is
    not writable on the hosting platform."""
    for d in (os.path.join(APP_DIR, "checkpoints"), os.path.join(tempfile.gettempdir(), "app_checkpoints")):
        try:
            os.makedirs(d, exist_ok=True)
            probe = os.path.join(d, ".write_test")
            open(probe, "w").close(); os.remove(probe)
            return d
        except OSError:
            continue
    return None


def ensure_remote_checkpoints(entries=None, token=None, use_api=None):
    """Download every checkpoint listed in remote_checkpoints.py that is not
    present yet (streamed to disk in 1 MB chunks, size-checked, written to a
    temporary name first so a broken download never looks like a valid
    file). Already-present files are skipped. Never raises.
    Returns (download_dir, messages)."""
    entries = REMOTE_CHECKPOINTS if entries is None else entries
    entries = [e for e in (entries or []) if e.get("filename") and e.get("url")]
    if not entries:
        return None, []
    dest = _download_dir()
    if dest is None:
        return None, ["Could not find a writable folder to store downloaded checkpoints."]
    token = token if token is not None else _github_token()
    msgs = []
    for e in entries:
        name, url = e["filename"].strip(), e["url"].strip()
        if not (CKPT_REG_RE.search(name) or CKPT_P2_RE.search(name)):
            msgs.append(f"`{name}` in remote_checkpoints.py does not follow the checkpoint naming pattern "
                        f"(checkpoint_reg_*.pkl / checkpoint_p2_<crtype>_*.pkl) — skipped.")
            continue
        path = os.path.join(dest, name)
        if os.path.exists(path) and os.path.getsize(path) > 0:
            continue
        tmp = path + ".part"
        try:
            with _open_remote(url, token, use_api) as resp, open(tmp, "wb") as f:
                expected = int(resp.headers.get("Content-Length") or 0)
                while True:
                    chunk = resp.read(1 << 20)
                    if not chunk:
                        break
                    f.write(chunk)
            got = os.path.getsize(tmp)
            if expected and got != expected:
                raise IOError(f"incomplete download ({got} of {expected} bytes)")
            with open(tmp, "rb") as f:
                if f.read(1) != b"\x80":  # every pickle starts with the PROTO opcode
                    raise ValueError("the downloaded file is not a pickle -- for a PRIVATE repository add "
                                     "GITHUB_TOKEN to the app's Secrets (see README)")
            os.replace(tmp, path)
            msgs.append(f"Downloaded `{name}` ({got / 1e6:.1f} MB) from its GitHub Release.")
        except Exception as ex:
            if os.path.exists(tmp):
                os.remove(tmp)
            hint = ""
            if "404" in str(ex) and GITHUB_RELEASE_RE.match(url) and not token:
                hint = " If the repository is PRIVATE, add GITHUB_TOKEN to the app's Secrets (see README)."
            msgs.append(f"Could not download `{name}`: {ex}.{hint}")
    return dest, msgs


def auto_load_once(st, base_dirs=DEFAULT_SEARCH_DIRS):
    """Called from app.py on every page but works ONCE per session: scan for
    checkpoints and load the default pair into EMPTY slots only (never
    overrides a model already loaded). Never raises."""
    if st.session_state.get("_auto_discover_tried"):
        return
    st.session_state["_auto_discover_tried"] = True
    msgs, warns = [], []
    dl_dir, dl_msgs = ensure_remote_checkpoints()
    msgs += dl_msgs
    if dl_dir and all(os.path.realpath(dl_dir) != os.path.realpath(b) for b in base_dirs if b):
        base_dirs = tuple(base_dirs) + (dl_dir,)
    try:
        cands = discover_checkpoint_candidates(base_dirs)
    except Exception as e:
        cands = []
        msgs.append(f"Scan failed: {e}")
    st.session_state["_ckpt_candidates"] = cands
    reg_c, clf_c = choose_default_pair(cands)
    if clf_c is not None and st.session_state.get("clf_res") is None:
        try:
            kind, obj, w = load_candidate(clf_c)
            if kind == "classification":
                activate_clf(st, obj, clf_c["path"], clf_c["name"], clf_c["source"])
                msgs.append(f"**Classification** model loaded from `{clf_c['path']}`.")
                warns += w
        except Exception as e:
            msgs.append(f"Could not load `{clf_c['path']}`: {e}")
    if reg_c is not None and st.session_state.get("reg_store") is None:
        try:
            kind, obj, w = load_candidate(reg_c)
            if kind == "regression":
                activate_reg(st, obj, reg_c["path"], reg_c["name"])
                msgs.append(f"**Regression** model loaded from `{reg_c['path']}` "
                            f"(default scenario: {st.session_state['reg_scenario']}).")
                warns += w
        except Exception as e:
            msgs.append(f"Could not load `{reg_c['path']}`: {e}")
    st.session_state["_auto_discover_msgs"] = msgs
    st.session_state["_load_warnings"] = warns


def auto_discover_checkpoints(base_dirs=DEFAULT_SEARCH_DIRS):
    """Kept for compatibility: load the default pair without a session."""
    cands = discover_checkpoint_candidates(base_dirs)
    reg_c, clf_c = choose_default_pair(cands)
    out = {"reg_store": None, "reg_path": None, "clf_res": None, "clf_path": None}
    for c, key, pkey, kind in ((reg_c, "reg_store", "reg_path", "regression"),
                               (clf_c, "clf_res", "clf_path", "classification")):
        if c is None:
            continue
        try:
            k, obj, _w = load_candidate(c)
            if k == kind:
                out[key], out[pkey] = obj, c["path"]
        except Exception:
            pass
    return out


# =====================================================================
# Predictions
# =====================================================================
def predict_cr(res, df_input, model_name=None):
    """df_input: raw app columns. Returns (df_input + Predicted_CR, Xt) using
    the chosen model (default: best) -- same logic as pred_all in E1."""
    name = resolve_model(res, model_name)
    model = res["fitted"][name][0]
    num_cols, cat_cols = extract_schema(res)
    Xt = _transform(res, df_input[num_cols + cat_cols])
    pred = model.predict(Xt)
    pred_raw = np.clip(np.expm1(pred) if res.get("use_log") else pred, 0, None)
    out = df_input.copy()
    out["Predicted_CR"] = pred_raw
    return out, Xt


def predict_risk(res, df_input, model_name=None):
    """df_input: raw app columns (numeric includes 'cr'). Returns df_input +
    Predicted_Risk + Confidence + Prob_<class> + Runner_Up + Margin, and Xt
    -- same logic as B2 in Experiment 3."""
    name = resolve_model(res, model_name)
    model = res["fitted"][name][0]
    num_cols, cat_cols = extract_schema(res)
    Xt = _transform(res, df_input[num_cols + cat_cols])
    proba = model.predict_proba(Xt)
    pred_idx = model.predict(Xt)
    out = df_input.copy()
    out["Predicted_Risk"] = [RISK_ORDER[int(i)] for i in pred_idx]
    out["Confidence"] = proba.max(axis=1)
    present = list(res.get("present", list(range(proba.shape[1]))))
    for j, cls_idx in enumerate(present):
        out[f"Prob_{RISK_ORDER[cls_idx]}"] = proba[:, j]
    if proba.shape[1] > 1:
        order = np.argsort(proba, axis=1)
        srt = np.sort(proba, axis=1)
        out["Runner_Up"] = [RISK_ORDER[present[j]] for j in order[:, -2]]
        out["Margin"] = srt[:, -1] - srt[:, -2]
    return out, Xt


# Backwards-compatible names used by earlier versions.
predict_regresi = predict_cr
predict_risiko = predict_risk


# =====================================================================
# Corrosion-rate reliability
# =====================================================================
K_NEIGHBORS = 5


def _to_raw(res, arr):
    arr = np.asarray(arr, dtype=float)
    return np.clip(np.expm1(arr) if res.get("use_log") else arr, 0, None)


def _reliability_context(res, model_name):
    """Pre-computed (and cached inside res) ingredients for CR reliability:
      1) test-set residuals of the chosen model in log1p space (for an
         empirical 80% prediction interval, conditional on prediction size),
      2) a reference set of REAL equipment (Xall_t + actual CR + identity
         columns, falling back to the training matrix) standardised for a
         nearest-neighbour search,
      3) the distribution of neighbour distances inside the reference set
         (for the applicability-domain check)."""
    from sklearn.neighbors import NearestNeighbors

    cache = res.setdefault("_rel_cache", {})
    name = resolve_model(res, model_name)
    if name in cache:
        return cache[name]
    ctx = {"name": name}

    # 1) test residuals
    yte, pte = res.get("yte_raw"), None
    preds = res.get("preds") or {}
    if name in preds and isinstance(preds[name], (tuple, list)) and len(preds[name]) == 2:
        pte = np.asarray(preds[name][1], dtype=float)
    elif res.get("Xte_t") is not None:
        try:
            Xte = pd.DataFrame(np.asarray(res["Xte_t"]), columns=res["feat_names"])
            pte = _to_raw(res, res["fitted"][name][0].predict(Xte))
        except Exception:
            pte = None
    if yte is not None and pte is not None and len(pte) == len(yte):
        y = np.asarray(yte, dtype=float)
        ok = np.isfinite(y) & np.isfinite(pte) & (y >= 0)
        if ok.sum() >= 10:
            p = np.clip(pte[ok], 0, None)
            r = np.log1p(y[ok]) - np.log1p(p)

            def q80(v):
                # Conservative (split-conformal style) order statistics: with
                # small samples, linearly interpolated quantiles systematically
                # give intervals that are too narrow (e.g. 9/13 = 69% instead
                # of 80%); 'lower'/'higher' guarantee >= 80% on the test set.
                return np.array([np.quantile(v, 0.10, method="lower"),
                                 np.quantile(v, 0.90, method="higher")])

            q_glob = q80(r)
            edges, bins = None, []
            if ok.sum() >= 60:  # condition on prediction size only with >= 20 errors per third
                edges = np.quantile(p, [1 / 3, 2 / 3])
                idx = np.digitize(p, edges)
                for b in range(3):
                    rb = r[idx == b]
                    bins.append(q80(rb) if len(rb) >= 20 else q_glob)
            ctx.update(edges=edges, bins=bins, q_global=q_glob, n_test=int(ok.sum()),
                       test_mae=float(np.mean(np.abs(y[ok] - p))))

    # 2) reference set of real equipment
    feat = res["feat_names"]
    ref_X = ref_y = ref_id = None
    if res.get("Xall_t") is not None and res.get("y_all_raw") is not None:
        ref_X = pd.DataFrame(np.asarray(res["Xall_t"]), columns=feat)
        ref_y = np.asarray(res["y_all_raw"], dtype=float)
        ida = res.get("identity_all")
        ref_id = ida.reset_index(drop=True) if isinstance(ida, pd.DataFrame) and len(ida) == len(ref_X) and ida.shape[1] else None
        ctx["ref_source"] = "real equipment records"
    elif res.get("Xtr_t") is not None and res.get("ytr_raw") is not None:
        ref_X = pd.DataFrame(np.asarray(res["Xtr_t"]), columns=feat)
        ref_y = np.asarray(res["ytr_raw"], dtype=float)
        ctx["ref_source"] = "training rows"
    if ref_X is not None and len(ref_X) > K_NEIGHBORS + 1 and len(ref_y) == len(ref_X):
        keep = np.isfinite(ref_y)
        ref_X, ref_y = ref_X[keep].reset_index(drop=True), ref_y[keep]
        if ref_id is not None:
            ref_id = ref_id[keep].reset_index(drop=True)
        if len(ref_X) > K_NEIGHBORS + 1:
            mu = ref_X.mean()
            sd = ref_X.std().replace(0, 1).fillna(1)
            Z = ((ref_X - mu) / sd).values
            nn = NearestNeighbors(n_neighbors=K_NEIGHBORS + 1).fit(Z)
            d, _ = nn.kneighbors(Z)
            own = np.sort(d[:, 1:].mean(axis=1))  # leave-one-out mean neighbour distance
            ctx.update(ref_X=ref_X, ref_y=ref_y, ref_id=ref_id, mu=mu, sd=sd, nn=nn, own=own)
    cache[name] = ctx
    return ctx


def _fmt_id(v):
    """Identity value as text: empty cells -> '', 109.0 -> '109'."""
    if v is None or (isinstance(v, float) and np.isnan(v)) or v is pd.NA:
        return ""
    if isinstance(v, (float, np.floating)) and float(v).is_integer():
        return str(int(v))
    return str(v).strip()


def _id_labels(ref_id):
    """'No | Equipment/Pipe Name | Equipment Part' per row. Robust to empty
    cells (the dataset has a row with an empty 'No'; with pandas >= 3
    astype(str) keeps NaN as a float and ' | '.join crashed)."""
    return np.array([" | ".join(p for p in (_fmt_id(v) for v in row) if p)
                     for row in ref_id.itertuples(index=False)], dtype=object)


def cr_reliability(res, Xt, pred, model_name=None):
    """Per-row reliability indicators for corrosion-rate predictions.

      CR_P10 / CR_P90      80% empirical prediction interval from the chosen
                           model's test-set errors (log scale, conditional on
                           the prediction size).
      Model_Agreement_CV   spread of ALL models of the scenario (std / mean).
      Domain_Percentile    distance of the input to its nearest real
                           equipment, as a percentile of typical distances
                           among real equipment -> Domain_Status Inside
                           (<= 90), Borderline (<= 99) or Outside.
      Similar_Median_CR / Similar_Min_CR / Similar_Max_CR   actual CR of the
                           K most similar real equipment.
      Within_Similar_Range prediction inside [min, max] of those neighbours.
      Most_Similar         identity of the closest real equipment.
      Confidence_Level     High / Medium / Low and Confidence_Reasons.
    Rules: Outside -> Low. Otherwise score = +1 Inside (0 Borderline)
    +1 agreement CV <= 0.15 (0 if <= 0.35, -1 above) +1 within neighbour
    range; score >= 2 High, >= 0 Medium, else Low."""
    ctx = _reliability_context(res, model_name)
    pred = np.asarray(pred, dtype=float)
    n = len(pred)
    out = pd.DataFrame(index=range(n))

    if "q_global" in ctx:
        lo_q = np.full(n, ctx["q_global"][0], dtype=float)
        hi_q = np.full(n, ctx["q_global"][1], dtype=float)
        if ctx.get("edges") is not None:
            idx = np.digitize(pred, ctx["edges"])
            for b in range(3):
                lo_q[idx == b], hi_q[idx == b] = ctx["bins"][b]
        out["CR_P10"] = np.clip(np.expm1(np.log1p(pred) + lo_q), 0, None)
        out["CR_P90"] = np.clip(np.expm1(np.log1p(pred) + hi_q), 0, None)
    else:
        out["CR_P10"] = np.nan
        out["CR_P90"] = np.nan

    all_preds = []
    for m in res.get("fitted", {}):
        try:
            all_preds.append(_to_raw(res, res["fitted"][m][0].predict(Xt)))
        except Exception:
            pass
    if len(all_preds) >= 2:
        A = np.vstack(all_preds)
        out["Model_Agreement_CV"] = A.std(axis=0) / np.maximum(A.mean(axis=0), 1e-6)
    else:
        out["Model_Agreement_CV"] = np.nan

    if "nn" in ctx:
        Z = ((Xt[ctx["ref_X"].columns] - ctx["mu"]) / ctx["sd"]).values
        d, idx = ctx["nn"].kneighbors(Z, n_neighbors=K_NEIGHBORS)
        md = d.mean(axis=1)
        out["Domain_Percentile"] = np.searchsorted(ctx["own"], md, side="right") / len(ctx["own"]) * 100
        out["Domain_Status"] = np.where(out["Domain_Percentile"] <= 90, "Inside",
                                        np.where(out["Domain_Percentile"] <= 99, "Borderline", "Outside"))
        ny = ctx["ref_y"][idx]
        out["Similar_Median_CR"] = np.median(ny, axis=1)
        out["Similar_Min_CR"] = ny.min(axis=1)
        out["Similar_Max_CR"] = ny.max(axis=1)
        out["Within_Similar_Range"] = (pred >= out["Similar_Min_CR"].values) & (pred <= out["Similar_Max_CR"].values)
        if ctx.get("ref_id") is not None:
            ids = _id_labels(ctx["ref_id"])
            out["Most_Similar"] = ids[idx[:, 0]]
        else:
            out["Most_Similar"] = [f"reference row #{i}" for i in idx[:, 0]]
    else:
        for c in ("Domain_Percentile", "Similar_Median_CR", "Similar_Min_CR", "Similar_Max_CR"):
            out[c] = np.nan
        out["Domain_Status"] = "Unknown"
        out["Within_Similar_Range"] = None
        out["Most_Similar"] = ""

    levels, reasons = [], []
    for i in range(n):
        dom = out["Domain_Status"].iloc[i]
        cv = out["Model_Agreement_CV"].iloc[i]
        within = out["Within_Similar_Range"].iloc[i]
        if dom == "Outside":
            levels.append("Low")
            reasons.append("far from known equipment (extrapolation)")
            continue
        score, why = 0, []
        if dom == "Inside":
            score += 1; why.append("similar to known equipment")
        elif dom == "Borderline":
            why.append("partly similar to known equipment")
        if pd.notna(cv):
            if cv <= 0.15:
                score += 1; why.append("models agree closely")
            elif cv <= 0.35:
                why.append("models partly agree")
            else:
                score -= 1; why.append("models disagree")
        if within is not None and not (isinstance(within, float) and np.isnan(within)):
            if bool(within):
                score += 1; why.append("CR within similar equipment range")
            else:
                why.append("CR outside similar equipment range")
        levels.append("High" if score >= 2 else ("Medium" if score >= 0 else "Low"))
        reasons.append("; ".join(why))
    out["Confidence_Level"] = levels
    out["Confidence_Reasons"] = reasons
    return out


def similar_equipment(res, Xt_row, model_name=None, overrides=None):
    """The K most similar REAL equipment to one input row: identity, actual
    CR, distance and their numeric feature values (original units)."""
    ctx = _reliability_context(res, model_name)
    if "nn" not in ctx:
        return None
    Z = ((Xt_row[ctx["ref_X"].columns] - ctx["mu"]) / ctx["sd"]).values
    d, idx = ctx["nn"].kneighbors(Z, n_neighbors=K_NEIGHBORS)
    num_cols, _ = extract_schema(res)
    rows = []
    for rank, (j, dist) in enumerate(zip(idx[0], d[0]), start=1):
        r = {"Rank": rank}
        if ctx.get("ref_id") is not None:
            r.update({str(k): _fmt_id(v) for k, v in ctx["ref_id"].iloc[j].to_dict().items()})
        r["Actual CR (mm/year)"] = float(ctx["ref_y"][j])
        r["Distance (std. units)"] = round(float(dist), 3)
        for i, c in enumerate(num_cols):
            r[resolve_display(c, overrides)] = float(ctx["ref_X"].iloc[j, i])
        rows.append(r)
    return pd.DataFrame(rows)


def reliability_basis(res, model_name=None):
    """One-line description of the data behind the reliability indicators."""
    ctx = _reliability_context(res, model_name)
    parts = []
    if "n_test" in ctx:
        parts.append(f"prediction interval from {ctx['n_test']} test-set errors of {ctx['name']}")
    if "nn" in ctx:
        parts.append(f"similarity search over {len(ctx['ref_X'])} {ctx['ref_source']}")
    parts.append(f"agreement across {len(res.get('fitted', {}))} models")
    return "; ".join(parts)


# =====================================================================
# SHAP
# =====================================================================
def _shap_3d(explainer, Xt):
    """SHAP values as (rows, features[, classes]) regardless of shap version
    (older versions return a list with one array per class)."""
    raw = explainer.shap_values(Xt)
    if isinstance(raw, list):
        return np.stack([np.asarray(a) for a in raw], axis=-1)
    arr = np.asarray(raw)
    if arr.ndim == 3 and arr.shape[0] != len(Xt) and arr.shape[1] == len(Xt):
        arr = np.moveaxis(arr, 0, -1)  # (classes, rows, features) -> (rows, features, classes)
    return arr


def compute_shap(res, Xt, model_name=None, how="signed"):
    """SHAP per original column (one-hot columns summed), mirroring D3 (Exp.
    1) / D6 (Exp. 2). For Stacking: mean over its tree-based base learners.

    how="abs"    -> |SHAP| per row, for multi-class averaged over classes
                   AFTER taking absolute values -- exactly D6
                   (np.abs(sv).mean(axis=(0, 2))). Use .mean() over rows for
                   feature importance.
    how="signed" -> signed SHAP; for multi-class the contribution towards
                   each row's PREDICTED class (why the model chose it).

    FIX: earlier versions averaged the SIGNED multi-class SHAP values over
    the classes. Class probabilities sum to 1, so those contributions cancel
    out and the result was ~1e-16 (numerical noise) for every feature."""
    name = resolve_model(res, model_name)
    base_names = [n for n in res["fitted"] if n != "Stacking"] if name == "Stacking" else [name]
    pred_labels = None
    sv_list = []
    for bn in base_names:
        mdl = res["fitted"][bn][0]
        try:
            arr = _shap_3d(shap.TreeExplainer(mdl), Xt)
        except Exception:
            continue
        if arr.ndim == 3:
            if how == "abs":
                arr = np.abs(arr).mean(axis=2)
            else:
                if pred_labels is None:
                    pred_labels = res["fitted"][name][0].predict(Xt)
                classes = list(getattr(mdl, "classes_", range(arr.shape[2])))
                idx = np.array([classes.index(l) if l in classes else 0 for l in pred_labels])
                arr = arr[np.arange(arr.shape[0]), :, idx]
        elif how == "abs":
            arr = np.abs(arr)
        sv_list.append(arr)
    if not sv_list:
        return None
    df_sv = pd.DataFrame(np.mean(sv_list, axis=0), columns=Xt.columns)
    return df_sv.T.groupby(df_sv.columns.map(base_column)).sum().T


# =====================================================================
# Excel input / output
# =====================================================================
# Result-column names shared by the UI and the Excel writer.
COL_ROW, COL_FILE = "#", "File Order"
COL_CR, COL_CR_CONF, COL_CR_WHY = "Predicted CR (mm/year)", "CR Confidence", "CR Confidence Reason"
COL_RISK, COL_RISK_CONF, COL_RISK_LVL, COL_RISK_WHY = ("Predicted Risk", "Risk Confidence",
                                                      "Risk Confidence Level", "Risk Confidence Note")
PCT_COLS = {COL_RISK_CONF, "Chance Low", "Chance Medium", "Chance Med-High", "Chance High",
            "Margin to Runner-up", "Model Spread"}
CR_VALUE_COLS = {COL_CR, "CR 80% Low", "CR 80% High", "Similar Equipment Median CR",
                 "Similar Equipment Min CR", "Similar Equipment Max CR", "Actual CR (mm/year)"}


def risk_conf_level(p):
    """Rule of thumb for the probability of the predicted risk class."""
    if p is None or pd.isna(p):
        return "Unknown"
    return "High" if p >= 0.70 else ("Medium" if p >= 0.50 else "Low")


def _hex_fill(hex_color):
    from openpyxl.styles import PatternFill
    h = hex_color.lstrip("#").upper()
    return PatternFill(start_color=h, end_color=h, fill_type="solid")


def xlsx_bytes(sheets):
    """{sheet name: DataFrame} -> .xlsx bytes. Auto column widths, header
    row frozen together with the row number / identity columns, coloured
    CR-confidence / risk cells, % and 4-decimal number formats."""
    from openpyxl.styles import Font
    buf = io.BytesIO()
    white_bold = Font(color="FFFFFF", bold=True)
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            if df is None:
                continue
            sheet = str(name)[:31]
            df.to_excel(xw, sheet_name=sheet, index=False)
            ws = xw.sheets[sheet]
            cols = list(df.columns)
            for i, col in enumerate(cols, start=1):
                letter = ws.cell(row=1, column=i).column_letter
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200).tolist()]) + 2
                ws.column_dimensions[letter].width = min(width, 45)
                colors = None
                if col in (COL_CR_CONF, COL_RISK_LVL):
                    colors = CONF_COLORS
                elif col == COL_RISK:
                    colors = RISK_COLORS
                fmt = "0.0%" if col in PCT_COLS else ("0.0000" if col in CR_VALUE_COLS else None)
                for r, v in enumerate(df[col].tolist(), start=2):
                    cell = ws.cell(row=r, column=i)
                    if colors and v in colors:
                        cell.fill = _hex_fill(colors[v]); cell.font = white_bold
                    elif col == COL_RISK_CONF and pd.notna(v):
                        cell.fill = _hex_fill(CONF_COLORS[risk_conf_level(v)]); cell.font = white_bold
                    if fmt:
                        cell.number_format = fmt
            lead = [c for c in (COL_CR, COL_RISK) if c in cols]
            if lead:  # freeze header + everything left of the first prediction column
                ws.freeze_panes = ws.cell(row=2, column=cols.index(lead[0]) + 1)
            elif len(df):
                ws.freeze_panes = "A2"
    return buf.getvalue()


def df_to_excel_bytes(df):  # backwards-compatible
    return xlsx_bytes({"Results": df})


def build_template_xlsx(specs):
    """Batch template. specs: list of dicts (col, header, label, kind,
    default, allowed). Sheets: 'Input' (headers only -- fill this one),
    'Column guide', 'Example' (training defaults) and 'Instructions'."""
    guide = pd.DataFrame([{
        "Excel header": s["header"], "Meaning": s["label"], "Internal column": s["col"],
        "Type": s["kind"], "Allowed values / training range": s["allowed"],
        "Default (used if left empty)": s["default"],
    } for s in specs])
    example = pd.DataFrame([{s["header"]: s["default"] for s in specs}])
    notes = pd.DataFrame({"How to use": [
        "Fill one row per equipment in the 'Input' sheet (keep the header row unchanged).",
        "Empty cells are allowed: they are filled with the default shown in 'Column guide'.",
        "Extra columns (e.g. Equipment ID, Tag No) are allowed and are kept in the results file.",
        "Categorical values should match the allowed values in 'Column guide'; unknown values are ignored by the model.",
        "Upload the file on the same page where this template was downloaded.",
    ]})
    return xlsx_bytes({"Input": pd.DataFrame(columns=[s["header"] for s in specs]),
                       "Column guide": guide, "Example": example, "Instructions": notes})


def excel_template_bytes(columns, display_map):  # backwards-compatible
    specs = [{"col": c, "header": display_map.get(c, c), "label": display_map.get(c, c),
              "kind": "", "default": "", "allowed": ""} for c in columns]
    return build_template_xlsx(specs)


def read_batch_file(uploaded, specs):
    """Read an uploaded .xlsx/.csv and map its headers to internal columns.
    A column is recognised by its Excel header, display label or internal
    name (case/space-insensitive). Reads the 'Input' sheet if present.
    Returns (df, missing_specs)."""
    name = getattr(uploaded, "name", "file").lower()
    if name.endswith(".csv"):
        df = pd.read_csv(uploaded)
    else:
        xl = pd.ExcelFile(uploaded)
        sheet = "Input" if "Input" in xl.sheet_names else xl.sheet_names[0]
        df = xl.parse(sheet)
    df = df.dropna(how="all")
    file_order = np.arange(1, len(df) + 1)

    def norm(s):
        return re.sub(r"\s+", " ", str(s)).strip().lower()

    lookup = {}
    for s in specs:
        for key in (s["header"], s["label"], s["col"]):
            lookup.setdefault(norm(key), s["col"])
    rename, used = {}, set()
    for c in df.columns:
        tgt = lookup.get(norm(c))
        if tgt and tgt not in used:
            rename[c] = tgt
            used.add(tgt)
    df = df.rename(columns=rename).reset_index(drop=True)
    df.attrs["file_order"] = file_order
    missing = [s for s in specs if s["col"] not in df.columns]
    return df, missing
